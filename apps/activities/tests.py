"""Testy plánu činností."""

from __future__ import annotations

import json
from datetime import date, datetime, timedelta
from pathlib import Path

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.test import Client, SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone

from apps.accounts.roles import ROLE_EMPLOYEE
from apps.employees.models import Employee, Employment
from apps.shifts.models import Shift, ShiftType
from apps.workplaces.models import EmployeeWorkplace, Workplace

from .models import ActivityItem, ActivityLocation, ActivityType
from .services import (
    ActivityConflictError,
    create_activity_item,
    employees_for_activity_day,
    locations_geojson_collection,
    save_activity_location,
)


SAMPLE_POLYGON = {
    "type": "Polygon",
    "coordinates": [
        [
            [14.90, 50.41],
            [14.91, 50.41],
            [14.91, 50.412],
            [14.90, 50.412],
            [14.90, 50.41],
        ]
    ],
}

SAMPLE_POLYGON_B = {
    "type": "Polygon",
    "coordinates": [
        [
            [14.902, 50.411],
            [14.905, 50.411],
            [14.905, 50.413],
            [14.902, 50.413],
            [14.902, 50.411],
        ]
    ],
}


class ActivityModuleTests(TestCase):
    def setUp(self):
        self.client = Client()
        User = get_user_model()
        self.staff = User.objects.create_user(
            username="act_admin",
            password="adminpass1",
            is_staff=True,
            is_superuser=True,
        )
        self.worker_user = User.objects.create_user(
            username="act_petr",
            password="petrpass1",
            is_staff=False,
        )
        Group.objects.get_or_create(name=ROLE_EMPLOYEE)
        self.worker_user.groups.add(Group.objects.get(name=ROLE_EMPLOYEE))
        self.employee = Employee.objects.create(
            first_name="Petr",
            last_name="Cinnost",
            internal_number="ACT001",
            user=self.worker_user,
        )
        self.workplace = Workplace.objects.create(
            name="Parkoviste Cinnost",
            address="Test 2",
        )
        EmployeeWorkplace.objects.create(
            employee=self.employee,
            workplace=self.workplace,
            valid_from=date(2026, 1, 1),
        )
        Employment.objects.create(
            employee=self.employee,
            started_on=date(2026, 1, 1),
            is_active=True,
        )
        self.type_r = ShiftType.objects.get(code="R")
        self.act_type = ActivityType.objects.get(code="kontrola")
        self.tz = timezone.get_current_timezone()
        self.day = date(2026, 10, 6)
        start = timezone.make_aware(
            datetime.combine(self.day, self.type_r.start_time), self.tz
        )
        end = timezone.make_aware(
            datetime.combine(self.day, self.type_r.end_time), self.tz
        )
        self.shift = Shift.objects.create(
            employee=self.employee,
            workplace=self.workplace,
            shift_type=self.type_r,
            starts_at=start,
            ends_at=end,
            status=Shift.Status.PUBLISHED,
        )
        self.main = ActivityLocation(
            code="SEV",
            name="Sever",
            geojson=json.dumps(SAMPLE_POLYGON),
            is_active=True,
            sort_order=5,
        )
        save_activity_location(user=self.staff, location=self.main, create=True)
        self.location = ActivityLocation(
            code="B2",
            name="Blok B2",
            description="Kontrola bloků B2.",
            parent=self.main,
            workplace=self.workplace,
            geojson=json.dumps(SAMPLE_POLYGON_B),
            is_active=True,
            sort_order=10,
        )
        save_activity_location(
            user=self.staff, location=self.location, create=True
        )

    def test_location_geojson_stored(self):
        loc = ActivityLocation.objects.get(code="B2")
        self.assertEqual(loc.geometry_dict()["type"], "Polygon")
        self.assertIn("coordinates", loc.geometry_dict())
        self.assertEqual(loc.parent_id, self.main.pk)
        self.assertIn("SEV ›", loc.hierarchy_label())

    def test_locations_geojson_collection_detail_props(self):
        collection = locations_geojson_collection()
        by_code = {
            f["properties"]["code"]: f["properties"]
            for f in collection["features"]
        }
        props = by_code["B2"]
        self.assertEqual(props["description"], "Kontrola bloků B2.")
        self.assertEqual(props["parent_code"], "SEV")
        self.assertEqual(props["workplace"], "Parkoviste Cinnost")
        self.assertIn("color", props)
        self.assertRegex(props["color"], r"^#[0-9A-Fa-f]{6}$")
        self.assertEqual(
            props["edit_url"],
            reverse("activity_location_edit", args=[self.location.pk]),
        )

    def test_location_create_map_shows_existing_context(self):
        self.client.login(username="act_admin", password="adminpass1")
        response = self.client.get(reverse("activity_location_create"))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("locations-context-geojson", html)
        self.assertIn("SEV", html)
        self.assertIn("B2", html)

        excluded = locations_geojson_collection(exclude_id=self.location.pk)
        codes = {f["properties"]["code"] for f in excluded["features"]}
        self.assertIn("SEV", codes)
        self.assertNotIn("B2", codes)

        edit = self.client.get(
            reverse("activity_location_edit", args=[self.location.pk])
        )
        self.assertEqual(edit.status_code, 200)
        ehtml = edit.content.decode()
        self.assertIn("locations-context-geojson", ehtml)
        self.assertIn("SEV", ehtml)

    def test_location_color_validation(self):
        loc = ActivityLocation(
            code="BAD",
            name="Spatna barva",
            geojson=json.dumps(SAMPLE_POLYGON),
            color="red",
        )
        with self.assertRaises(ValidationError):
            loc.full_clean()
        self.main.color = "#237A35"
        save_activity_location(
            user=self.staff, location=self.main, create=False
        )
        self.main.refresh_from_db()
        self.assertEqual(self.main.color, "#237A35")

    def test_reject_nested_sublocation(self):
        deeper = ActivityLocation(
            code="B2A",
            name="Roh",
            parent=self.location,
            geojson=json.dumps(SAMPLE_POLYGON_B),
        )
        with self.assertRaises(ValidationError):
            deeper.full_clean()

    def test_activity_on_main_and_sub(self):
        on_main = create_activity_item(
            user=self.staff,
            employee=self.employee,
            shift=self.shift,
            location=self.main,
            activity_type=self.act_type,
            starts_at=self.shift.starts_at,
            ends_at=self.shift.starts_at + timedelta(minutes=30),
            sort_order=5,
        )
        on_sub = create_activity_item(
            user=self.staff,
            employee=self.employee,
            shift=self.shift,
            location=self.location,
            activity_type=self.act_type,
            starts_at=self.shift.starts_at + timedelta(minutes=30),
            ends_at=self.shift.starts_at + timedelta(hours=1),
            sort_order=10,
        )
        self.assertEqual(on_main.location_id, self.main.pk)
        self.assertEqual(on_sub.location_id, self.location.pk)
        self.assertEqual(ActivityItem.objects.count(), 2)

    def test_create_activity_within_published_shift(self):
        item = create_activity_item(
            user=self.staff,
            employee=self.employee,
            shift=self.shift,
            location=self.location,
            activity_type=self.act_type,
            starts_at=self.shift.starts_at,
            ends_at=self.shift.starts_at + timedelta(hours=1),
            note="Kontrola B2",
            sort_order=10,
        )
        self.assertEqual(item.day, self.day)
        self.assertEqual(item.location.code, "B2")
        self.assertEqual(ActivityItem.objects.count(), 1)

    def test_reject_outside_shift_window(self):
        with self.assertRaises(ActivityConflictError):
            create_activity_item(
                user=self.staff,
                employee=self.employee,
                shift=self.shift,
                location=self.location,
                activity_type=self.act_type,
                starts_at=self.shift.starts_at - timedelta(minutes=30),
                ends_at=self.shift.starts_at + timedelta(hours=1),
            )

    def test_export_ok_for_manager(self):
        create_activity_item(
            user=self.staff,
            employee=self.employee,
            shift=self.shift,
            location=self.location,
            activity_type=self.act_type,
            starts_at=self.shift.starts_at,
            ends_at=self.shift.starts_at + timedelta(hours=1),
            note="Zkontrolovat závory",
        )
        self.client.login(username="act_admin", password="adminpass1")
        url = (
            reverse("activity_export")
            + f"?employee={self.employee.pk}&day={self.day.isoformat()}"
        )
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Denní plán činností", content)
        self.assertIn("B2", content)
        self.assertIn("SEV", content)
        self.assertIn("Cinnost", content)
        self.assertIn("Zkontrolovat závory", content)
        self.assertIn("Činnosti (hodinové)", content)
        self.assertIn("Denní úkoly", content)

        xlsx_url = (
            reverse("activity_export_xlsx")
            + f"?employee={self.employee.pk}&day={self.day.isoformat()}"
        )
        xlsx = self.client.get(xlsx_url)
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn(
            "spreadsheetml.sheet",
            xlsx["Content-Type"],
        )
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(xlsx.content))
        self.assertIn("Cinnosti", wb.sheetnames)
        self.assertIn("Denni_ukoly", wb.sheetnames)
        self.assertIn("Zkontrolovat závory", str(wb["Cinnosti"].cell(2, 7).value))

        board = self.client.get(
            reverse("activity_day")
            + f"?day={self.day.isoformat()}&workplace={self.workplace.pk}"
        )
        self.assertEqual(board.status_code, 200)
        bhtml = board.content.decode()
        self.assertIn("activity-slot-item-type", bhtml)
        self.assertIn("Kontrola", bhtml)
        self.assertIn("data-activity-open-export", bhtml)
        self.assertIn("activity-export-dialog", bhtml)
        self.assertNotIn("Export Excel (pracoviště)", bhtml)
        self.assertNotIn("Export PDF (pracoviště)", bhtml)

    def test_export_requires_permission(self):
        User = get_user_model()
        from apps.accounts.roles import ROLE_MANAGER

        manager = User.objects.create_user(
            username="act_mgr_noexp",
            password="mgrpass1",
            is_staff=True,
            is_superuser=False,
        )
        mgr_group, _ = Group.objects.get_or_create(name=ROLE_MANAGER)
        manager.groups.add(mgr_group)
        # Explicitně bez exportu (matice role ho má — odebereme)
        from django.contrib.auth.models import Permission

        export_perm = Permission.objects.get(codename="export_activity_plan")
        mgr_group.permissions.remove(export_perm)
        manager.user_permissions.remove(export_perm)

        self.client.login(username="act_mgr_noexp", password="mgrpass1")
        export_url = (
            reverse("activity_export")
            + f"?employee={self.employee.pk}&day={self.day.isoformat()}"
        )
        denied = self.client.get(export_url)
        self.assertEqual(denied.status_code, 302)
        board = self.client.get(
            reverse("activity_day")
            + f"?day={self.day.isoformat()}&workplace={self.workplace.pk}"
        )
        self.assertEqual(board.status_code, 200)
        self.assertNotIn("data-activity-open-export", board.content.decode())

        manager.user_permissions.add(export_perm)
        # refresh permission cache
        self.client.logout()
        self.client.login(username="act_mgr_noexp", password="mgrpass1")
        allowed = self.client.get(export_url)
        self.assertEqual(allowed.status_code, 200)
        board2 = self.client.get(
            reverse("activity_day")
            + f"?day={self.day.isoformat()}&workplace={self.workplace.pk}"
        )
        self.assertIn("data-activity-open-export", board2.content.decode())

    def test_board_and_employee_detail_are_exclusive(self):
        self.client.login(username="act_admin", password="adminpass1")
        board_url = (
            reverse("activity_day")
            + f"?day={self.day.isoformat()}&workplace={self.workplace.pk}"
        )
        board = self.client.get(board_url)
        self.assertEqual(board.status_code, 200)
        board_html = board.content.decode()
        self.assertIn("Rozvrh ·", board_html)
        self.assertNotIn("Zpět na rozvrh", board_html)

        detail = self.client.get(
            board_url + f"&employee={self.employee.pk}"
        )
        self.assertEqual(detail.status_code, 200)
        detail_html = detail.content.decode()
        self.assertNotIn("Rozvrh ·", detail_html)
        self.assertIn("Cinnost Petr", detail_html)
        self.assertIn("Zpět na rozvrh", detail_html)
        self.assertIn(
            f'href="?day={self.day.isoformat()}&workplace={self.workplace.pk}"',
            detail_html,
        )

    def test_employee_cannot_plan(self):
        self.client.login(username="act_petr", password="petrpass1")
        response = self.client.get(reverse("activity_day"))
        self.assertEqual(response.status_code, 403)
        response2 = self.client.get(reverse("activity_location_list"))
        self.assertEqual(response2.status_code, 403)

    def test_filter_form_day_before_employee(self):
        from .forms import DayPlanFilterForm

        form = DayPlanFilterForm(employee_queryset=Employee.objects.all())
        self.assertEqual(
            list(form.fields.keys()), ["day", "workplace", "employee"]
        )
        self.assertIn("onchange", form.fields["day"].widget.attrs)

    def test_month_calendar_syncs_with_day(self):
        from .calendar_month import activity_month_markers

        self.client.login(username="act_admin", password="adminpass1")
        url = (
            reverse("activity_day")
            + f"?day={self.day.isoformat()}&workplace={self.workplace.pk}"
        )
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("activity-month-cal", html)
        self.assertIn("activity-cal-day--selected", html)
        self.assertIn(f'value="{self.day.isoformat()}"', html)
        self.assertIn(f"?day={self.day.isoformat()}", html)

        markers = activity_month_markers(
            self.staff, self.workplace, self.day.year, self.day.month
        )
        self.assertTrue(markers.get(self.day, {}).get("has_shift"))

        create_activity_item(
            user=self.staff,
            employee=self.employee,
            shift=self.shift,
            location=self.main,
            activity_type=self.act_type,
            starts_at=self.shift.starts_at,
            ends_at=self.shift.starts_at + timedelta(hours=1),
        )
        markers2 = activity_month_markers(
            self.staff, self.workplace, self.day.year, self.day.month
        )
        self.assertTrue(markers2[self.day]["has_plan"])
        page2 = self.client.get(url)
        self.assertContains(page2, "activity-cal-day--planned")

    def test_day_filter_requires_published_shift_that_day(self):
        future = Employee.objects.create(
            first_name="Novy",
            last_name="Nastup",
            internal_number="ACT002",
        )
        EmployeeWorkplace.objects.create(
            employee=future,
            workplace=self.workplace,
            valid_from=date(2026, 11, 1),
        )
        Employment.objects.create(
            employee=future,
            started_on=date(2026, 11, 1),
            is_active=True,
        )
        nov_day = date(2026, 11, 2)
        start = timezone.make_aware(
            datetime.combine(nov_day, self.type_r.start_time), self.tz
        )
        end = timezone.make_aware(
            datetime.combine(nov_day, self.type_r.end_time), self.tz
        )
        Shift.objects.create(
            employee=future,
            workplace=self.workplace,
            shift_type=self.type_r,
            starts_at=start,
            ends_at=end,
            status=Shift.Status.PUBLISHED,
        )

        early = date(2026, 9, 30)
        qs_early = employees_for_activity_day(self.staff, early)
        self.assertNotIn(future, qs_early)
        self.assertIn(self.employee, employees_for_activity_day(self.staff, self.day))

        self.client.login(username="act_admin", password="adminpass1")
        bad = self.client.get(
            reverse("activity_day")
            + f"?employee={future.pk}&day={early.isoformat()}"
        )
        self.assertEqual(bad.status_code, 200)
        self.assertContains(bad, "publikovanou směnu práce")
        self.assertNotContains(bad, "Novy Nastup · 30.09.2026")

        export = self.client.get(
            reverse("activity_export")
            + f"?employee={future.pk}&day={early.isoformat()}"
        )
        self.assertEqual(export.status_code, 302)

        ok = self.client.get(
            reverse("activity_day")
            + f"?employee={future.pk}&day={nov_day.isoformat()}"
        )
        self.assertEqual(ok.status_code, 200)
        self.assertContains(ok, "Nastup")
        self.assertIn(
            future, employees_for_activity_day(self.staff, nov_day)
        )

    def test_employment_without_shift_not_listed(self):
        """Nástup sám o sobě nestačí — bez publikované směny práce není v seznamu."""
        starter = Employee.objects.create(
            first_name="Bez",
            last_name="Smeny",
            internal_number="ACT003",
        )
        EmployeeWorkplace.objects.create(
            employee=starter,
            workplace=self.workplace,
            valid_from=date(2026, 11, 1),
        )
        Employment.objects.create(
            employee=starter,
            started_on=date(2026, 11, 1),
            is_active=True,
        )
        nov_day = date(2026, 11, 2)
        self.assertNotIn(
            starter, employees_for_activity_day(self.staff, nov_day)
        )
        self.client.login(username="act_admin", password="adminpass1")
        response = self.client.get(
            reverse("activity_day")
            + f"?day={nov_day.isoformat()}&workplace={self.workplace.pk}"
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "publikovanou směnu práce")
        self.assertContains(response, "Samotný nástup")


class ActivityScheduleTests(TestCase):
    def setUp(self):
        from apps.activities.tests import ActivityModuleTests

        # reuse fixture pattern
        ActivityModuleTests.setUp(self)

    def test_allocate_hours_2_to_1(self):
        from .schedule import allocate_hours_by_weights

        self.assertEqual(allocate_hours_by_weights(6, [2, 1]), [4, 2])
        self.assertEqual(sum(allocate_hours_by_weights(7, [2, 1])), 7)

    def test_hour_range_label(self):
        from .schedule import hour_range_label

        start = timezone.make_aware(
            datetime.combine(self.day, datetime.strptime("08:00", "%H:%M").time()),
            self.tz,
        )
        end = start + timedelta(hours=1)
        self.assertEqual(hour_range_label(start, end), "8–9")
        partial_end = start + timedelta(hours=1, minutes=30)
        # 08:00–09:30
        self.assertEqual(hour_range_label(start, partial_end), "8–9:30")

    def test_slot_form_prefills_od_do(self):
        from datetime import time as dt_time

        self.client.login(username="act_admin", password="adminpass1")
        start = timezone.make_aware(
            datetime.combine(self.day, self.type_r.start_time), self.tz
        )
        end = start + timedelta(hours=1)
        url = (
            reverse("activity_slot_form")
            + f"?employee={self.employee.pk}&day={self.day.isoformat()}"
            f"&starts_at={start.strftime('%Y-%m-%dT%H:%M')}"
            f"&ends_at={end.strftime('%Y-%m-%dT%H:%M')}"
        )
        empty = self.client.get(url)
        self.assertEqual(empty.status_code, 200)
        html = empty.content.decode()
        self.assertIn('data-slot-start="08:00"', html)
        self.assertIn('data-slot-end="09:00"', html)
        self.assertIn('value="08:00"', html)
        self.assertIn('value="09:00"', html)
        self.assertIn("activity-slot-note", html)
        self.assertIn("<textarea", html)

        from .schedule import replace_slot_items

        replace_slot_items(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            starts_at=start,
            ends_at=end,
            rows=[
                {
                    "location": None,
                    "activity_type": self.act_type,
                    "note": "A",
                    "time_from": dt_time(8, 0),
                    "time_to": dt_time(8, 20),
                }
            ],
        )
        filled = self.client.get(url)
        self.assertEqual(filled.status_code, 200)
        fhtml = filled.content.decode()
        self.assertIn('value="08:00"', fhtml)
        self.assertIn('value="08:20"', fhtml)

    def test_generate_from_preferences(self):
        from .models import ActivityPreference
        from .schedule import generate_day_plan

        ActivityPreference.objects.create(
            employee=self.employee,
            location=self.main,
            frequency=2,
            sort_order=10,
            default_activity_type=self.act_type,
        )
        ActivityPreference.objects.create(
            employee=self.employee,
            location=self.location,
            frequency=1,
            sort_order=20,
            default_activity_type=self.act_type,
        )
        items = generate_day_plan(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            replace_existing=True,
        )
        self.assertGreater(len(items), 0)
        minutes_main = sum(
            int((i.ends_at - i.starts_at).total_seconds() // 60)
            for i in items
            if i.location_id == self.main.pk
        )
        minutes_sub = sum(
            int((i.ends_at - i.starts_at).total_seconds() // 60)
            for i in items
            if i.location_id == self.location.pk
        )
        # směna R typicky 6 h → 4 h + 2 h
        self.assertEqual(minutes_main, 240)
        self.assertEqual(minutes_sub, 120)

    def test_preferences_page_allows_many_rows(self):
        self.client.login(username="act_admin", password="adminpass1")
        url = reverse("activity_preferences", args=[self.employee.pk])
        response = self.client.get(url + f"?day={self.day.isoformat()}")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("Šablona lokalit", html)
        self.assertIn("Přidat lokalitu", html)
        self.assertIn("data-pref-max=\"20\"", html)
        self.assertIn("activity_preferences.js", html)

    def test_generate_without_prefs_fails(self):
        from .schedule import generate_day_plan

        with self.assertRaises(ActivityConflictError):
            generate_day_plan(
                user=self.staff,
                employee=self.employee,
                day=self.day,
                replace_existing=True,
            )

    def test_slot_multiple_items_and_optional_location(self):
        from .schedule import replace_slot_items
        from django.utils import timezone as tz

        start = tz.make_aware(
            datetime.combine(self.day, self.type_r.start_time), self.tz
        )
        end = start + timedelta(hours=1)
        created = replace_slot_items(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            starts_at=start,
            ends_at=end,
            rows=[
                {
                    "location": self.main,
                    "activity_type": self.act_type,
                    "note": "Kontrola",
                },
                {
                    "location": None,
                    "activity_type": self.act_type,
                    "note": "Telefonát",
                },
            ],
        )
        self.assertEqual(len(created), 2)
        self.assertIsNone(created[1].location_id)
        self.assertEqual(created[1].note, "Telefonát")

        replace_slot_items(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            starts_at=start,
            ends_at=end,
            rows=[],
        )
        self.assertEqual(
            ActivityItem.objects.filter(
                employee=self.employee, day=self.day
            ).count(),
            0,
        )

    def test_slot_max_ten(self):
        from .schedule import replace_slot_items
        from django.utils import timezone as tz

        start = tz.make_aware(
            datetime.combine(self.day, self.type_r.start_time), self.tz
        )
        end = start + timedelta(hours=1)
        rows = [
            {
                "location": None,
                "activity_type": self.act_type,
                "note": f"r{i}",
            }
            for i in range(11)
        ]
        with self.assertRaises(ActivityConflictError):
            replace_slot_items(
                user=self.staff,
                employee=self.employee,
                day=self.day,
                starts_at=start,
                ends_at=end,
                rows=rows,
            )

    def test_slot_partial_times_inside_hour(self):
        from datetime import time as dt_time

        from .schedule import hour_items_for_day, replace_slot_items
        from django.utils import timezone as tz

        start = tz.make_aware(
            datetime.combine(self.day, self.type_r.start_time), self.tz
        )
        end = start + timedelta(hours=1)
        created = replace_slot_items(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            starts_at=start,
            ends_at=end,
            rows=[
                {
                    "location": None,
                    "activity_type": self.act_type,
                    "note": "Krátká",
                    "time_from": dt_time(8, 10),
                    "time_to": dt_time(8, 30),
                }
            ],
        )
        self.assertEqual(len(created), 1)
        local_s = timezone.localtime(created[0].starts_at)
        local_e = timezone.localtime(created[0].ends_at)
        self.assertEqual(local_s.strftime("%H:%M"), "08:10")
        self.assertEqual(local_e.strftime("%H:%M"), "08:30")
        self.assertFalse(created[0].all_day)

        with self.assertRaises(ActivityConflictError):
            replace_slot_items(
                user=self.staff,
                employee=self.employee,
                day=self.day,
                starts_at=start,
                ends_at=end,
                rows=[
                    {
                        "location": None,
                        "activity_type": self.act_type,
                        "note": "Mimo",
                        "time_from": dt_time(7, 0),
                        "time_to": dt_time(8, 30),
                    }
                ],
            )

        full = replace_slot_items(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            starts_at=start,
            ends_at=end,
            rows=[
                {
                    "location": None,
                    "activity_type": self.act_type,
                    "note": "Celá",
                }
            ],
        )
        self.assertEqual(
            timezone.localtime(full[0].starts_at).strftime("%H:%M"),
            timezone.localtime(start).strftime("%H:%M"),
        )
        self.assertEqual(len(hour_items_for_day(self.employee, self.day)), 1)

    def test_day_task_outside_hour_grid(self):
        from .schedule import (
            create_day_task,
            generate_day_plan,
            hour_items_for_day,
        )
        from .services import day_tasks_for_day
        from .models import ActivityPreference

        task = create_day_task(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            location=self.main,
            activity_type=self.act_type,
            note="B2 kontrola 4×",
        )
        self.assertTrue(task.all_day)
        self.assertEqual(day_tasks_for_day(self.employee, self.day), [task])
        self.assertEqual(hour_items_for_day(self.employee, self.day), [])

        ActivityPreference.objects.create(
            employee=self.employee,
            location=self.main,
            frequency=1,
            sort_order=10,
        )
        generated = generate_day_plan(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            replace_existing=True,
        )
        self.assertTrue(all(not g.all_day for g in generated))
        self.assertEqual(
            ActivityItem.objects.filter(
                employee=self.employee, day=self.day, all_day=True
            ).count(),
            1,
        )
        self.assertGreater(len(hour_items_for_day(self.employee, self.day)), 0)

        self.client.login(username="act_admin", password="adminpass1")
        page = self.client.get(
            reverse("activity_day")
            + f"?day={self.day.isoformat()}&employee={self.employee.pk}"
            f"&workplace={self.workplace.pk}"
        )
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Denní úkoly")
        self.assertContains(page, "B2 kontrola 4×")

    def test_day_task_long_note(self):
        from .schedule import create_day_task

        long_note = "Kontrola B2\n" + ("řádek " * 40) + "\nKonec."
        self.assertGreater(len(long_note), 255)
        task = create_day_task(
            user=self.staff,
            employee=self.employee,
            day=self.day,
            location=None,
            activity_type=self.act_type,
            note=long_note,
        )
        task.refresh_from_db()
        self.assertEqual(task.note, long_note.strip())


class MapTilesPolicyTests(SimpleTestCase):
    def test_js_avoids_blocked_osm_tiles(self):
        root = Path(__file__).resolve().parents[2] / "static" / "mpdmb" / "js"
        for name in (
            "activity_map_tiles.js",
            "activity_locations_map.js",
            "activity_location_editor.js",
        ):
            text = (root / name).read_text(encoding="utf-8")
            self.assertNotIn("tile.openstreetmap.org/", text)
            self.assertNotIn("basemaps.cartocdn.com", text)
            self.assertNotIn("tile.openstreetmap.fr/osmfr/", text)
        tiles = (root / "activity_map_tiles.js").read_text(encoding="utf-8")
        self.assertIn("tile.openstreetmap.de/", tiles)

    def test_overview_map_keeps_city_center(self):
        """Přehled lokalit nesmí přepsat střed města slepým fitBounds."""
        root = Path(__file__).resolve().parents[2] / "static" / "mpdmb" / "js"
        text = (root / "activity_locations_map.js").read_text(encoding="utf-8")
        self.assertNotIn(".fitBounds(", text)
        self.assertIn("setView", text)


class MapCityCenterLocaleTests(TestCase):
    def setUp(self):
        self.client = Client()
        User = get_user_model()
        self.staff = User.objects.create_user(
            username="map_admin",
            password="adminpass1",
            is_staff=True,
            is_superuser=True,
        )

    def test_location_list_center_attrs_use_dot_decimal(self):
        self.client.login(username="map_admin", password="adminpass1")
        response = self.client.get(reverse("activity_location_list") + "?city=MB")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn('data-center-lat="50.411350"', html)
        self.assertIn('data-center-lon="14.903180"', html)
        self.assertNotIn("data-center-lat=\"50,411", html)
        self.assertNotIn("data-center-lon=\"14,903", html)
