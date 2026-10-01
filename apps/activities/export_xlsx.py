"""Export plánu činností do Excelu (openpyxl)."""

from __future__ import annotations

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .export_pack import ActivityExportPack

HEADER_FILL = PatternFill("solid", fgColor="166999")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="E8F2F8")


def _safe_text(value) -> str:
    text = "" if value is None else str(value)
    if text[:1] in ("=", "+", "-", "@"):
        return f"'{text}"
    return text


def _style_header(ws, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_activity_workbook(pack: ActivityExportPack) -> bytes:
    wb = Workbook()

    ws_sum = wb.active
    ws_sum.title = "Souhrn"
    ws_sum.append(
        [
            "Den",
            "Rozsah",
            "Pracoviště",
            "Počet osob",
            "Činností (hodinových)",
            "Denních úkolů",
            "Vytvořeno",
        ]
    )
    _style_header(ws_sum, 7)
    hour_n = sum(len(p.hour_items) for p in pack.persons)
    day_n = sum(len(p.day_tasks) for p in pack.persons)
    ws_sum.append(
        [
            pack.day.strftime("%d.%m.%Y"),
            _safe_text(pack.scope_label),
            _safe_text(pack.workplace.name if pack.workplace else ""),
            len(pack.persons),
            hour_n,
            day_n,
            pack.exported_at.strftime("%d.%m.%Y %H:%M"),
        ]
    )
    ws_sum.freeze_panes = "A2"
    _autosize(ws_sum, [12, 28, 22, 12, 18, 14, 18])

    ws_act = wb.create_sheet("Cinnosti")
    act_headers = [
        "Zaměstnanec",
        "Od",
        "Do",
        "Typ",
        "Lokalita",
        "Vozidlo",
        "Očekávání",
        "Pořadí",
    ]
    ws_act.append(act_headers)
    _style_header(ws_act, len(act_headers))
    row_i = 0
    for person in pack.persons:
        emp_label = f"{person.employee.last_name} {person.employee.first_name}"
        for it in person.hour_items:
            row_i += 1
            values = [
                _safe_text(emp_label),
                timezone.localtime(it.starts_at).strftime("%H:%M"),
                timezone.localtime(it.ends_at).strftime("%H:%M"),
                _safe_text(it.activity_type.name),
                _safe_text(
                    it.location.hierarchy_label() if it.location_id else ""
                ),
                _safe_text(it.vehicle.plate if it.vehicle_id else ""),
                _safe_text(it.note),
                it.sort_order,
            ]
            ws_act.append(values)
            if row_i % 2 == 0:
                for col in range(1, len(act_headers) + 1):
                    ws_act.cell(row_i + 1, col).fill = ALT_FILL
            ws_act.cell(row_i + 1, 7).alignment = Alignment(wrap_text=True)
    ws_act.freeze_panes = "A2"
    ws_act.auto_filter.ref = ws_act.dimensions
    _autosize(ws_act, [22, 8, 8, 16, 28, 12, 40, 10])

    ws_day = wb.create_sheet("Denni_ukoly")
    day_headers = ["Zaměstnanec", "Typ", "Lokalita", "Vozidlo", "Poznámka", "Pořadí"]
    ws_day.append(day_headers)
    _style_header(ws_day, len(day_headers))
    row_i = 0
    for person in pack.persons:
        emp_label = f"{person.employee.last_name} {person.employee.first_name}"
        for it in person.day_tasks:
            row_i += 1
            ws_day.append(
                [
                    _safe_text(emp_label),
                    _safe_text(it.activity_type.name),
                    _safe_text(
                        it.location.hierarchy_label() if it.location_id else ""
                    ),
                    _safe_text(it.vehicle.plate if it.vehicle_id else ""),
                    _safe_text(it.note),
                    it.sort_order,
                ]
            )
            if row_i % 2 == 0:
                for col in range(1, len(day_headers) + 1):
                    ws_day.cell(row_i + 1, col).fill = ALT_FILL
            ws_day.cell(row_i + 1, 5).alignment = Alignment(wrap_text=True)
    ws_day.freeze_panes = "A2"
    ws_day.auto_filter.ref = ws_day.dimensions
    _autosize(ws_day, [22, 16, 28, 12, 50, 10])

    ws_loc = wb.create_sheet("Lokality")
    loc_headers = ["Kód", "Název", "Hierarchie", "Popis", "Barva"]
    ws_loc.append(loc_headers)
    _style_header(ws_loc, len(loc_headers))
    for i, loc in enumerate(pack.locations):
        ws_loc.append(
            [
                _safe_text(loc.code),
                _safe_text(loc.name),
                _safe_text(loc.hierarchy_label()),
                _safe_text(loc.description),
                _safe_text(loc.color),
            ]
        )
        if i % 2 == 1:
            for col in range(1, len(loc_headers) + 1):
                ws_loc.cell(i + 2, col).fill = ALT_FILL
    ws_loc.freeze_panes = "A2"
    _autosize(ws_loc, [10, 20, 28, 40, 10])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
