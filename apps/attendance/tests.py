"""Testy docházky z publikovaného plánu směn."""

from datetime import date, datetime, time

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from apps.accounts.roles import ROLE_EMPLOYEE
from apps.employees.models import Employee, Employment
from apps.shifts.models import Shift, ShiftType
from apps.workplaces.models import EmployeeWorkplace, Workplace

from .models import Absence, AbsenceType, WorkInterval
from .services import (
    confirm_work_interval,
    month_attendance_matrix,
    set_absence_for_shift,
)


class AttendanceFromPlanTests(TestCase):
    def setUp(self):
        self.client = Client()
        User = get_user_model()
        self.staff = User.objects.create_user(
            username="att_admin",
            password="adminpass1",
            is_staff=True,
            is_superuser=True,
        )
        self.worker_user = User.objects.create_user(
            username="att_petr",
            password="petrpass1",
            is_staff=False,
        )
        Group.objects.get_or_create(name=ROLE_EMPLOYEE)
        self.worker_user.groups.add(Group.objects.get(name=ROLE_EMPLOYEE))
        self.employee = Employee.objects.create(
            first_name="Petr",
            last_name="Dochazka",
            internal_number="ATT001",
            user=self.worker_user,
        )
        self.other = Employee.objects.create(
            first_name="Jana",
            last_name="Cizi",
            internal_number="ATT002",
        )
        self.workplace = Workplace.objects.create(
            name="Parkoviste Test",
            address="Test 1",
        )
        EmployeeWorkplace.objects.create(
            employee=self.employee,
            workplace=self.workplace,
            valid_from=date(2026, 1, 1),
        )
        EmployeeWorkplace.objects.create(
            employee=self.other,
            workplace=self.workplace,
            valid_from=date(2026, 1, 1),
        )
        Employment.objects.create(
            employee=self.employee,
            started_on=date(2026, 1, 1),
            is_active=True,
        )
        Employment.objects.create(
            employee=self.other,
            started_on=date(2026, 1, 1),
            is_active=True,
        )
        self.type_r = ShiftType.objects.get(code="R")
        self.tz = timezone.get_current_timezone()
        self.day = date(2026, 10, 5)
        start = timezone.make_aware(
            datetime.combine(self.day, self.type_r.start_time), self.tz
        )
        end = timezone.make_aware(
            datetime.combine(self.day, self.type_r.end_time), self.tz
        )
        self.shift = Shift.objects.create(
            employee=self.employee,
            workplace=self.workplace,
            shift_type=self.type_r,
            starts_at=start,
            ends_at=end,
            status=Shift.Status.PUBLISHED,
        )
        self.other_shift = Shift.objects.create(
            employee=self.other,
            workplace=self.workplace,
            shift_type=self.type_r,
            starts_at=start,
            ends_at=end,
            status=Shift.Status.PUBLISHED,
        )

    def test_confirm_shift_creates_work_interval(self):
        interval = confirm_work_interval(user=self.staff, shift=self.shift)
        self.assertEqual(interval.status, WorkInterval.Status.CONFIRMED)
        self.assertEqual(interval.starts_at, self.shift.starts_at)
        self.assertEqual(interval.ends_at, self.shift.ends_at)
        self.assertEqual(interval.shift_id, self.shift.pk)
        self.assertFalse(Absence.objects.filter(shift=self.shift).exists())

    def test_absence_sick_uses_planned_net_no_work_interval(self):
        sick = AbsenceType.objects.get(code="nemoc")
        absence = set_absence_for_shift(
            user=self.staff, shift=self.shift, absence_type=sick
        )
        self.assertEqual(absence.planned_missed_minutes, self.shift.planned_net_minutes())
        self.assertEqual(absence.absence_type.code, "nemoc")
        self.assertFalse(
            WorkInterval.objects.filter(shift=self.shift).exists()
        )

    def test_unresolved_in_month_summary(self):
        matrix = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        self.assertGreaterEqual(matrix["summary"]["unresolved_count"], 1)
        row = next(
            r for r in matrix["rows"] if r.employee.pk == self.employee.pk
        )
        cell = next(c for c in row.cells if c.day == self.day)
        self.assertEqual(cell.status, "unresolved")

        confirm_work_interval(user=self.staff, shift=self.shift)
        matrix2 = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row2 = next(
            r for r in matrix2["rows"] if r.employee.pk == self.employee.pk
        )
        cell2 = next(c for c in row2.cells if c.day == self.day)
        self.assertEqual(cell2.status, "work")

    def test_employee_cannot_see_other_attendance(self):
        self.client.login(username="att_petr", password="petrpass1")
        url = reverse("attendance_month") + "?year=2026&month=10"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Dochazka", content)
        self.assertNotIn("Cizi", content)

        day_url = (
            reverse("attendance_day")
            + f"?employee={self.other.pk}&day=2026-10-05&year=2026&month=10"
        )
        denied = self.client.get(day_url)
        self.assertEqual(denied.status_code, 403)

    def test_employee_can_confirm_own_shift(self):
        self.client.login(username="att_petr", password="petrpass1")
        url = reverse("attendance_confirm", args=[self.shift.pk])
        response = self.client.post(
            url,
            {"year": "2026", "month": "10", "as_plan": "1"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            WorkInterval.objects.filter(
                shift=self.shift, status=WorkInterval.Status.CONFIRMED
            ).exists()
        )

    def test_attendance_row_shows_workplace_segment_label(self):
        matrix = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row = next(
            r for r in matrix["rows"] if r.employee.pk == self.employee.pk
        )
        self.assertEqual(row.workplace_name, self.workplace.name)
        self.assertIn(self.workplace.name, row.workplace_segments_label)
        self.assertIn("od 01.01.2026", row.workplace_segments_label)

    def test_cannot_confirm_and_absence_together(self):
        confirm_work_interval(user=self.staff, shift=self.shift)
        sick = AbsenceType.objects.get(code="nemoc")
        with self.assertRaises(Exception):
            set_absence_for_shift(
                user=self.staff, shift=self.shift, absence_type=sick
            )

    def test_employee_without_employment_in_month_hidden(self):
        future = Employee.objects.create(
            first_name="Stanislav",
            last_name="HolubTest",
            internal_number="ATT003",
        )
        EmployeeWorkplace.objects.create(
            employee=future,
            workplace=self.workplace,
            valid_from=date(2026, 10, 1),
        )
        Employment.objects.create(
            employee=future,
            started_on=date(2026, 11, 1),
            is_active=True,
        )
        matrix_sep = month_attendance_matrix(
            user=self.staff, year=2026, month=9
        )
        ids_sep = {r.employee.pk for r in matrix_sep["rows"]}
        self.assertNotIn(future.pk, ids_sep)

        oct_day = date(2026, 10, 6)
        start = timezone.make_aware(
            datetime.combine(oct_day, self.type_r.start_time), self.tz
        )
        end = timezone.make_aware(
            datetime.combine(oct_day, self.type_r.end_time), self.tz
        )
        Shift.objects.create(
            employee=future,
            workplace=self.workplace,
            shift_type=self.type_r,
            starts_at=start,
            ends_at=end,
            status=Shift.Status.PUBLISHED,
        )
        matrix_oct = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        ids_oct = {r.employee.pk for r in matrix_oct["rows"]}
        self.assertIn(future.pk, ids_oct)

        matrix_nov = month_attendance_matrix(
            user=self.staff, year=2026, month=11
        )
        ids_nov = {r.employee.pk for r in matrix_nov["rows"]}
        self.assertIn(future.pk, ids_nov)
        self.assertIn(self.employee.pk, ids_nov)

    def test_ad_hoc_work_on_saturday(self):
        saturday = date(2026, 10, 10)  # sobota
        self.assertEqual(saturday.weekday(), 5)
        tz = self.tz
        start = timezone.make_aware(
            datetime.combine(saturday, time(9, 0)), tz
        )
        end = timezone.make_aware(
            datetime.combine(saturday, time(13, 0)), tz
        )
        from .services import create_ad_hoc_work_interval

        interval = create_ad_hoc_work_interval(
            user=self.staff,
            employee=self.employee,
            workplace=self.workplace,
            starts_at=start,
            ends_at=end,
            break_minutes=0,
        )
        self.assertIsNone(interval.shift_id)
        self.assertEqual(interval.status, WorkInterval.Status.CONFIRMED)
        matrix = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row = next(
            r for r in matrix["rows"] if r.employee.pk == self.employee.pk
        )
        cell = next(c for c in row.cells if c.day == saturday)
        self.assertEqual(cell.status, "work")
        self.assertGreater(cell.worked_minutes, 0)

    def test_employee_cannot_create_ad_hoc(self):
        self.client.login(username="att_petr", password="petrpass1")
        saturday = date(2026, 10, 10)
        url = (
            reverse("attendance_adhoc_form")
            + f"?employee={self.employee.pk}&day={saturday.isoformat()}"
            f"&year=2026&month=10"
        )
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_manager_adhoc_form_uses_time_fields(self):
        self.client.login(username="att_admin", password="adminpass1")
        saturday = date(2026, 10, 10)
        url = (
            reverse("attendance_adhoc_form")
            + f"?employee={self.employee.pk}&day={saturday.isoformat()}"
            f"&year=2026&month=10"
        )
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn('type="time"', html)
        self.assertNotIn("datetime-local", html)
        self.assertContains(response, saturday.strftime("%d.%m.%Y"))

        post = self.client.post(
            url,
            {
                "employee": self.employee.pk,
                "day": saturday.isoformat(),
                "year": 2026,
                "month": 10,
                "workplace": self.workplace.pk,
                "start_time": "09:00",
                "end_time": "13:00",
                "break_minutes": 0,
                "note": "",
            },
        )
        self.assertEqual(post.status_code, 302)
        from .models import WorkInterval

        wi = WorkInterval.objects.filter(
            employment__employee=self.employee, shift__isnull=True
        ).latest("pk")
        self.assertEqual(timezone.localtime(wi.starts_at).time(), time(9, 0))
        self.assertEqual(timezone.localtime(wi.ends_at).time(), time(13, 0))

    def test_month_empty_cell_has_context_menu_markup(self):
        self.client.login(username="att_admin", password="adminpass1")
        response = self.client.get(
            reverse("attendance_month") + "?year=2026&month=10"
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "att-empty-menu")
        self.assertContains(response, "att-adhoc-dialog")
        self.assertContains(response, "att-adhoc-type")
        self.assertContains(response, "attendance_cell_menu.js")

    def test_adhoc_partial_shift_type_creates_work(self):
        self.client.login(username="att_admin", password="adminpass1")
        saturday = date(2026, 10, 10)
        url = reverse("attendance_adhoc_form")
        response = self.client.post(
            url,
            {
                "partial": "1",
                "employee": self.employee.pk,
                "day": saturday.isoformat(),
                "year": 2026,
                "month": 10,
                "shift_type": self.type_r.pk,
                "workplace": self.workplace.pk,
            },
            HTTP_ACCEPT="application/json",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data.get("ok"), data)
        wi = WorkInterval.objects.filter(
            employment__employee=self.employee, shift__isnull=True
        ).latest("pk")
        self.assertEqual(timezone.localtime(wi.starts_at).time(), self.type_r.start_time)
        self.assertEqual(timezone.localtime(wi.ends_at).time(), self.type_r.end_time)
        self.assertEqual(wi.status, WorkInterval.Status.CONFIRMED)

    def test_month_long_plan_preview_rows(self):
        self.client.login(username="att_admin", password="adminpass1")
        response = self.client.get(
            reverse("attendance_month") + "?year=2026&month=10&plan=1"
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "att-plan-row")
        self.assertContains(response, "Dlouhodobý plán")
        self.assertContains(response, "attendance_plan_toggle.js")
        matrix = month_attendance_matrix(user=self.staff, year=2026, month=10)
        row = next(
            r for r in matrix["rows"] if r.employee.pk == self.employee.pk
        )
        cell = next(c for c in row.cells if c.day == self.day)
        self.assertTrue(cell.plan_codes)
        self.assertIn(self.type_r.code, cell.plan_codes)

    def test_confirm_form_partial_get_and_post(self):
        self.client.login(username="att_admin", password="adminpass1")
        url = reverse("attendance_confirm_form", args=[self.shift.pk])
        get_res = self.client.get(
            url + "?partial=1&year=2026&month=10"
        )
        self.assertEqual(get_res.status_code, 200)
        html = get_res.content.decode()
        self.assertNotIn("<html", html.lower())
        self.assertIn("att-confirm-form", html)
        self.assertIn('type="time"', html)
        self.assertIn("Upravit a potvrdit práci", html)

        post = self.client.post(
            url,
            {
                "partial": "1",
                "year": 2026,
                "month": 10,
                "start_time": "08:00",
                "end_time": "16:00",
                "break_minutes": 30,
                "note": "",
            },
            HTTP_ACCEPT="application/json",
        )
        self.assertEqual(post.status_code, 200)
        self.assertTrue(post.json().get("ok"), post.json())
        self.assertTrue(
            WorkInterval.objects.filter(
                shift=self.shift, status=WorkInterval.Status.CONFIRMED
            ).exists()
        )

    def test_month_has_confirm_dialog(self):
        self.client.login(username="att_admin", password="adminpass1")
        response = self.client.get(
            reverse("attendance_month") + "?year=2026&month=10"
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "att-confirm-dialog")
        self.assertContains(response, "att-confirm-dialog-host")
        self.assertContains(response, "att-day-dialog")
        self.assertContains(response, "att-day-dialog-host")

    def test_attendance_day_partial_fragment(self):
        self.client.login(username="att_admin", password="adminpass1")
        url = (
            reverse("attendance_day")
            + f"?employee={self.employee.pk}&day=2026-10-05"
            f"&year=2026&month=10&partial=1"
        )
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertNotIn("<html", html.lower())
        self.assertIn("att-day-modal", html)
        self.assertIn("Potvrdit jako plán", html)
        self.assertIn("partial", html)

    def test_confirm_without_employment_returns_link(self):
        from apps.attendance.services import AttendanceConflictError

        Employment.objects.filter(employee=self.employee).delete()
        with self.assertRaises(AttendanceConflictError) as ctx:
            confirm_work_interval(user=self.staff, shift=self.shift)
        self.assertEqual(ctx.exception.code, "missing_employment")
        self.assertIn("pracovní vztah", str(ctx.exception.messages[0]))

        self.client.login(username="att_admin", password="adminpass1")
        url = reverse("attendance_confirm", args=[self.shift.pk])
        response = self.client.post(
            url,
            {"partial": "1", "year": 2026, "month": 10, "as_plan": "1"},
            HTTP_ACCEPT="application/json",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data.get("ok"))
        self.assertIn("pracovní vztah", data.get("error", ""))
        self.assertIn(
            reverse("employee_edit", args=[self.employee.pk]),
            data.get("employee_url", ""),
        )

        month = self.client.get(
            reverse("attendance_month") + "?year=2026&month=10"
        )
        self.assertEqual(month.status_code, 200)
        self.assertContains(month, "Chybí pracovní vztah")
        self.assertContains(month, "data-missing-employment=\"1\"")

    def test_meal_vouchers_count_by_threshold(self):
        from apps.attendance.models import MealAllowanceSettings
        from apps.attendance.services.matrix import meal_vouchers_for_day

        settings = MealAllowanceSettings.load()
        settings.min_worked_minutes = 180
        settings.second_worked_minutes = 720
        settings.is_active = True
        settings.save()

        self.assertEqual(meal_vouchers_for_day(179, settings), 0)
        self.assertEqual(meal_vouchers_for_day(180, settings), 1)
        self.assertEqual(meal_vouchers_for_day(719, settings), 1)
        self.assertEqual(meal_vouchers_for_day(720, settings), 2)

        settings.second_worked_minutes = None
        settings.save()
        self.assertEqual(meal_vouchers_for_day(720, settings), 1)

        settings.second_worked_minutes = 720
        settings.is_active = False
        settings.save()
        self.assertEqual(meal_vouchers_for_day(720, settings), 0)

        settings.is_active = True
        settings.save()

        matrix0 = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row0 = next(
            r for r in matrix0["rows"] if r.employee.pk == self.employee.pk
        )
        self.assertEqual(row0.meal_vouchers, 0)

        confirm_work_interval(user=self.staff, shift=self.shift)
        matrix1 = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row1 = next(
            r for r in matrix1["rows"] if r.employee.pk == self.employee.pk
        )
        # směna R = 6 h → 1. stravenka, ne 2.
        self.assertEqual(row1.meal_vouchers, 1)
        self.assertEqual(matrix1["summary"]["meal_vouchers"], 1)
        self.assertEqual(matrix1["meal_threshold_minutes"], 180)
        self.assertEqual(matrix1["meal_second_threshold_minutes"], 720)

        start_long = timezone.make_aware(
            datetime.combine(self.day, time(6, 0)), self.tz
        )
        end_long = timezone.make_aware(
            datetime.combine(self.day, time(18, 0)), self.tz
        )
        confirm_work_interval(
            user=self.staff,
            shift=self.shift,
            starts_at=start_long,
            ends_at=end_long,
            break_minutes=0,
        )
        matrix_long = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row_long = next(
            r for r in matrix_long["rows"] if r.employee.pk == self.employee.pk
        )
        self.assertEqual(row_long.meal_vouchers, 2)

        settings.second_worked_minutes = None
        settings.save()
        matrix_one = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row_one = next(
            r for r in matrix_one["rows"] if r.employee.pk == self.employee.pk
        )
        self.assertEqual(row_one.meal_vouchers, 1)
        self.assertIsNone(matrix_one["meal_second_threshold_minutes"])

        settings.second_worked_minutes = 720
        settings.min_worked_minutes = 24 * 60
        settings.save()
        matrix2 = month_attendance_matrix(
            user=self.staff, year=2026, month=10
        )
        row2 = next(
            r for r in matrix2["rows"] if r.employee.pk == self.employee.pk
        )
        self.assertEqual(row2.meal_vouchers, 0)

        self.client.login(username="att_admin", password="adminpass1")
        page = self.client.get(
            reverse("attendance_month") + "?year=2026&month=10"
        )
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Stravné")
        self.assertNotContains(page, "<h2>Legenda</h2>")

    def test_attendance_xlsx_export_for_manager(self):
        from io import BytesIO

        from openpyxl import load_workbook

        confirm_work_interval(user=self.staff, shift=self.shift)
        self.client.login(username="att_admin", password="adminpass1")
        url = reverse("attendance_export_xlsx") + "?year=2026&month=10"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "spreadsheetml.sheet",
            response["Content-Type"],
        )
        self.assertIn("dochazka_2026-10.xlsx", response["Content-Disposition"])
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Souhrn", "Denni", "Metodika"])
        souhrn = wb["Souhrn"]
        self.assertEqual(souhrn["A1"].value, "Zaměstnanec")
        self.assertTrue(souhrn.max_row >= 2)

        self.client.login(username="att_petr", password="petrpass1")
        denied = self.client.get(url)
        self.assertEqual(denied.status_code, 403)


class MealAllowanceSettingsTests(TestCase):
    def test_load_creates_default(self):
        from apps.attendance.models import MealAllowanceSettings

        obj = MealAllowanceSettings.load()
        self.assertEqual(obj.pk, 1)
        self.assertEqual(obj.min_worked_minutes, 180)
        self.assertEqual(obj.second_worked_minutes, 720)
        self.assertTrue(obj.is_active)
        self.assertEqual(MealAllowanceSettings.objects.count(), 1)
        self.assertEqual(MealAllowanceSettings.load().pk, 1)

    def test_second_threshold_must_exceed_first(self):
        from django.core.exceptions import ValidationError

        from apps.attendance.models import MealAllowanceSettings

        obj = MealAllowanceSettings.load()
        obj.min_worked_minutes = 180
        obj.second_worked_minutes = 180
        with self.assertRaises(ValidationError):
            obj.full_clean()
        obj.second_worked_minutes = 720
        obj.full_clean()
        obj.second_worked_minutes = None
        obj.full_clean()
