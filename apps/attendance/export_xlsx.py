"""Export měsíční docházky do Excelu (openpyxl)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from django.contrib.auth.models import AbstractBaseUser
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.shifts.fund import format_hours

from .services import month_attendance_matrix

HEADER_FILL = PatternFill("solid", fgColor="166999")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="E8F2F8")


def _minutes_to_hours_decimal(minutes: int) -> Decimal:
    return (Decimal(minutes) / Decimal(60)).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )


def _safe_text(value) -> str:
    text = "" if value is None else str(value)
    if text[:1] in ("=", "+", "-", "@"):
        return f"'{text}"
    return text


def _style_header(ws, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_attendance_workbook(
    *, user: AbstractBaseUser, year: int, month: int
) -> bytes:
    matrix = month_attendance_matrix(user=user, year=year, month=month)
    rows = matrix["rows"]
    absence_types = matrix.get("absence_types") or []

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Souhrn"

    sum_headers = [
        "Zaměstnanec",
        "Interní číslo",
        "Pracoviště",
        "Plán (h:mm)",
        "Plán (h)",
        "Práce (h:mm)",
        "Práce (h)",
        "Absence (h:mm)",
        "Absence (h)",
        "Nevyřešeno",
        "Stravné",
    ]
    for t in absence_types:
        sum_headers.append(f"Absence {t.code}")
    ws_sum.append(sum_headers)
    _style_header(ws_sum, len(sum_headers))

    for i, row in enumerate(rows):
        emp = row.employee
        values = [
            _safe_text(f"{emp.last_name} {emp.first_name}"),
            _safe_text(emp.internal_number),
            _safe_text(row.workplace_name),
            _safe_text(row.planned_label),
            float(_minutes_to_hours_decimal(row.planned_minutes)),
            _safe_text(row.worked_label),
            float(_minutes_to_hours_decimal(row.worked_minutes)),
            _safe_text(row.missed_label),
            float(_minutes_to_hours_decimal(row.missed_minutes)),
            row.unresolved_count,
            row.meal_vouchers,
        ]
        for t in absence_types:
            values.append(row.absence_by_code.get(t.code, 0))
        ws_sum.append(values)
        if i % 2 == 1:
            for col in range(1, len(sum_headers) + 1):
                ws_sum.cell(i + 2, col).fill = ALT_FILL

    summary = matrix.get("summary") or {}
    total_row = [
        "CELKEM",
        "",
        "",
        _safe_text(summary.get("planned_label", format_hours(0))),
        float(_minutes_to_hours_decimal(summary.get("planned_minutes", 0))),
        _safe_text(summary.get("worked_label", format_hours(0))),
        float(_minutes_to_hours_decimal(summary.get("worked_minutes", 0))),
        _safe_text(summary.get("missed_label", format_hours(0))),
        float(_minutes_to_hours_decimal(summary.get("missed_minutes", 0))),
        summary.get("unresolved_count", 0),
        summary.get("meal_vouchers", 0),
    ]
    for _t in absence_types:
        total_row.append("")
    ws_sum.append(total_row)
    for col in range(1, len(sum_headers) + 1):
        ws_sum.cell(ws_sum.max_row, col).font = Font(bold=True)

    ws_sum.freeze_panes = "A2"
    ws_sum.auto_filter.ref = f"A1:{get_column_letter(len(sum_headers))}{max(1, ws_sum.max_row)}"
    _autosize(ws_sum, [22, 14, 18, 12, 10, 12, 10, 12, 10, 12, 10] + [12] * len(absence_types))

    ws_day = wb.create_sheet("Denni")
    day_headers = [
        "Zaměstnanec",
        "Interní číslo",
        "Datum",
        "Den",
        "Stav",
        "Plán kódy",
        "Plán (h:mm)",
        "Plán (h)",
        "Práce (h:mm)",
        "Práce (h)",
        "Absence (h:mm)",
        "Absence (h)",
        "Svátek",
    ]
    ws_day.append(day_headers)
    _style_header(ws_day, len(day_headers))

    day_idx = 0
    for row in rows:
        emp = row.employee
        name = f"{emp.last_name} {emp.first_name}"
        for cell in row.cells:
            if cell.status == "empty" and not cell.plan_codes:
                continue
            day_idx += 1
            values = [
                _safe_text(name),
                _safe_text(emp.internal_number),
                cell.day,
                _safe_text(cell.weekday_label),
                _safe_text(cell.status_label),
                _safe_text(cell.plan_codes),
                _safe_text(format_hours(cell.planned_minutes)),
                float(_minutes_to_hours_decimal(cell.planned_minutes)),
                _safe_text(format_hours(cell.worked_minutes)),
                float(_minutes_to_hours_decimal(cell.worked_minutes)),
                _safe_text(format_hours(cell.missed_minutes)),
                float(_minutes_to_hours_decimal(cell.missed_minutes)),
                _safe_text(cell.holiday_name or ""),
            ]
            ws_day.append(values)
            date_cell = ws_day.cell(ws_day.max_row, 3)
            date_cell.number_format = "DD.MM.YYYY"
            if day_idx % 2 == 0:
                for col in range(1, len(day_headers) + 1):
                    ws_day.cell(ws_day.max_row, col).fill = ALT_FILL

    ws_day.freeze_panes = "A2"
    if ws_day.max_row > 1:
        ws_day.auto_filter.ref = (
            f"A1:{get_column_letter(len(day_headers))}{ws_day.max_row}"
        )
    _autosize(ws_day, [22, 14, 12, 6, 14, 12, 12, 10, 12, 10, 12, 10, 18])

    meta = wb.create_sheet("Metodika")
    meta.append(["Pole", "Hodnota"])
    _style_header(meta, 2)
    meta_rows = [
        ("Období", f"{month:02d}.{year}"),
        ("Export", "Docházka (návrh, ne mzdová uzávěrka)"),
        ("Vytvořeno", date.today().strftime("%d.%m.%Y")),
        (
            "1. stravenka (min)",
            matrix.get("meal_threshold_minutes")
            if matrix.get("meal_threshold_minutes") is not None
            else "vypnuto",
        ),
        (
            "2. stravenka (min)",
            matrix.get("meal_second_threshold_minutes")
            if matrix.get("meal_second_threshold_minutes") is not None
            else "—",
        ),
        ("Desetinné hodiny", "7:30 = 7,50 h"),
    ]
    for label, value in meta_rows:
        meta.append([label, _safe_text(value)])
    _autosize(meta, [28, 40])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
