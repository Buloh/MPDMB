"""Práce s .docx: placeholdery, MERGEFIELD a vyplnění."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterable

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from .placeholders import PlaceholderDef, allowed_keys_for_scope, placeholders_for_scope

TOKEN_RE = re.compile(r"\{\{([a-z0-9_.]+)\}\}")
MERGEFIELD_RE = re.compile(
    r"MERGEFIELD\s+([a-zA-Z0-9_.]+)",
    re.IGNORECASE,
)
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def create_blank_docx_bytes() -> bytes:
    buffer = io.BytesIO()
    Document().save(buffer)
    return strip_mail_merge_settings(buffer.getvalue())


def strip_mail_merge_settings(docx_bytes: bytes) -> bytes:
    """Odstraní vazbu Korespondence (CSV/SQL dialog při otevření ve Wordu)."""
    source = io.BytesIO(docx_bytes)
    output = io.BytesIO()
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/settings.xml":
                root = etree.fromstring(data)
                for node in list(root):
                    if etree.QName(node).localname == "mailMerge":
                        root.remove(node)
                data = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            zout.writestr(info, data)
    return output.getvalue()


def _iter_paragraphs(document: Document) -> Iterable[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _iter_table_paragraphs(table)
    for section in document.sections:
        for header in (
            section.header,
            section.first_page_header,
            section.even_page_header,
        ):
            if header is None:
                continue
            yield from header.paragraphs
            for table in header.tables:
                yield from _iter_table_paragraphs(table)
        for footer in (
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            if footer is None:
                continue
            yield from footer.paragraphs
            for table in footer.tables:
                yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _comments_root(document: Document):
    """Vrátí kořen w:comments (vytvoří part, pokud chybí)."""
    part = document.part
    try:
        comments_part = part.part_related_by(RT.COMMENTS)
    except KeyError:
        comments_xml = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<w:comments xmlns:w="http://schemas.openxmlformats.org/'
            b'wordprocessingml/2006/main"></w:comments>'
        )
        comments_part = Part(
            PackURI("/word/comments.xml"),
            CT.WML_COMMENTS,
            comments_xml,
            part.package,
        )
        part.relate_to(comments_part, RT.COMMENTS)
    return etree.fromstring(comments_part.blob), comments_part


def _next_comment_id(comments_elm) -> int:
    max_id = -1
    for comment in comments_elm.findall(f"{W_NS}comment"):
        try:
            max_id = max(max_id, int(comment.get(qn("w:id"))))
        except (TypeError, ValueError):
            continue
    return max_id + 1


def _add_word_comment(
    document: Document,
    paragraph: Paragraph,
    help_text: str,
    *,
    author: str = "SYNERA",
) -> None:
    """Přidá komentář Wordu na konec odstavce (nápověda při najetí)."""
    if not (help_text or "").strip():
        return
    comments_elm, comments_part = _comments_root(document)
    cid = _next_comment_id(comments_elm)

    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), str(cid))
    paragraph._p.append(start)

    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), str(cid))
    paragraph._p.append(end)

    run = paragraph.add_run()
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "CommentReference")
    r_pr.append(r_style)
    run._r.insert(0, r_pr)
    ref = OxmlElement("w:commentReference")
    ref.set(qn("w:id"), str(cid))
    run._r.append(ref)

    comment = etree.SubElement(comments_elm, f"{W_NS}comment")
    comment.set(qn("w:id"), str(cid))
    comment.set(qn("w:author"), author)
    comment.set(qn("w:initials"), "SY")
    c_p = etree.SubElement(comment, f"{W_NS}p")
    c_r = etree.SubElement(c_p, f"{W_NS}r")
    c_t = etree.SubElement(c_r, f"{W_NS}t")
    c_t.text = help_text.strip()
    comments_part._blob = etree.tostring(
        comments_elm,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )


def _add_merge_field_to_paragraph(
    paragraph: Paragraph,
    field_name: str,
    *,
    display_label: str | None = None,
    document: Document | None = None,
    help_text: str = "",
) -> None:
    """Vloží MERGEFIELD; viditelný text = český popisek, volitelně komentář."""
    label = (display_label or field_name).strip() or field_name
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" MERGEFIELD {field_name} "
    run._r.append(instr)

    run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)

    run = paragraph.add_run()
    run.text = f"«{label}»"

    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)

    if document is not None and help_text:
        _add_word_comment(document, paragraph, help_text)


def create_universal_template_bytes(
    scope: str,
    *,
    title: str,
) -> bytes:
    """Univerzální .docx se všemi MERGEFIELD daného okruhu."""
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(
        "Univerzální šablona SYNERA. Pole můžete ve Wordu přesunout, "
        "smazat nepotřebná a upravit vzhled. Při generování v aplikaci "
        "se sloučená pole nahradí údaji z karty. Nápověda k poli je "
        "v komentáři Wordu (najetí na značku komentáře)."
    )
    document.add_heading("Seznam polí", level=2)
    for item in placeholders_for_scope(scope):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{item.label}: ")
        _add_merge_field_to_paragraph(
            paragraph,
            item.key,
            display_label=item.label,
            document=document,
            help_text=item.help_text or item.label,
        )
    buffer = io.BytesIO()
    document.save(buffer)
    return strip_mail_merge_settings(buffer.getvalue())


def load_document(source) -> Document:
    """Načte Document z bytes, cesty nebo souboru (UploadedFile nezavírá natrvalo)."""
    if isinstance(source, (bytes, bytearray)):
        return Document(io.BytesIO(source))
    if hasattr(source, "storage") and hasattr(source, "name") and source.name:
        with source.open("rb") as handle:
            return Document(io.BytesIO(handle.read()))
    if hasattr(source, "read") and not isinstance(source, (str,)):
        data = source.read()
        if hasattr(source, "seek"):
            try:
                source.seek(0)
            except Exception:
                pass
        return Document(io.BytesIO(data))
    return Document(str(source))


def _extract_mergefield_keys_from_element(element) -> list[str]:
    found: list[str] = []
    for node in element.iter():
        tag = node.tag
        if tag == f"{W_NS}instrText" and node.text:
            match = MERGEFIELD_RE.search(node.text)
            if match:
                found.append(match.group(1))
        elif tag == f"{W_NS}fldSimple":
            instr = node.get(qn("w:instr")) or ""
            match = MERGEFIELD_RE.search(instr)
            if match:
                found.append(match.group(1))
    return found


def extract_tokens_from_document(document: Document) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()

    def add(key: str) -> None:
        if key not in seen:
            seen.add(key)
            found.append(key)

    for paragraph in _iter_paragraphs(document):
        for match in TOKEN_RE.finditer(paragraph.text or ""):
            add(match.group(1))
        for key in _extract_mergefield_keys_from_element(paragraph._p):
            add(key)
    return found


def validate_placeholders_in_docx(source, scope: str) -> list[str]:
    """Vrátí seznam neznámých klíčů. Prázdný = OK."""
    document = load_document(source)
    allowed = allowed_keys_for_scope(scope)
    return [
        key
        for key in extract_tokens_from_document(document)
        if key not in allowed
    ]


def append_placeholder_token(
    source,
    key: str,
    *,
    display_label: str | None = None,
    help_text: str = "",
) -> bytes:
    """Připíše MERGEFIELD na konec dokumentu (nová verze šablony)."""
    document = load_document(source)
    paragraph = document.add_paragraph()
    _add_merge_field_to_paragraph(
        paragraph,
        key,
        display_label=display_label,
        document=document,
        help_text=help_text,
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _safe_text(value: str) -> str:
    text = value or ""
    if text[:1] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def _replace_mergefields_in_paragraph(
    paragraph: Paragraph,
    values: dict[str, str],
    missing: list[str],
    seen_missing: set[str],
) -> None:
    element = paragraph._p

    for fld in list(element.iter(f"{W_NS}fldSimple")):
        instr = fld.get(qn("w:instr")) or ""
        match = MERGEFIELD_RE.search(instr)
        if not match:
            continue
        key = match.group(1)
        raw = values.get(key, "")
        if key not in values or not raw:
            if key not in seen_missing:
                seen_missing.add(key)
                missing.append(key)
        text = _safe_text(raw if key in values else "")
        parent = fld.getparent()
        if parent is None:
            continue
        new_run = OxmlElement("w:r")
        new_text = OxmlElement("w:t")
        new_text.set(qn("xml:space"), "preserve")
        new_text.text = text
        new_run.append(new_text)
        parent.replace(fld, new_run)

    children = list(element)
    i = 0
    while i < len(children):
        child = children[i]
        begin = False
        if child.tag == f"{W_NS}r":
            for fld_char in child.findall(qn("w:fldChar")):
                if fld_char.get(qn("w:fldCharType")) == "begin":
                    begin = True
                    break
        if not begin:
            i += 1
            continue
        j = i
        instr_parts: list[str] = []
        end_idx = None
        while j < len(children):
            run_el = children[j]
            if run_el.tag == f"{W_NS}r":
                for instr in run_el.findall(qn("w:instrText")):
                    if instr.text:
                        instr_parts.append(instr.text)
                for fld_char in run_el.findall(qn("w:fldChar")):
                    if fld_char.get(qn("w:fldCharType")) == "end":
                        end_idx = j
                        break
            if end_idx is not None:
                break
            j += 1
        if end_idx is None:
            i += 1
            continue
        instr_text = "".join(instr_parts)
        match = MERGEFIELD_RE.search(instr_text)
        if not match:
            i = end_idx + 1
            continue
        key = match.group(1)
        raw = values.get(key, "")
        if key not in values or not raw:
            if key not in seen_missing:
                seen_missing.add(key)
                missing.append(key)
        text = _safe_text(raw if key in values else "")
        new_run = OxmlElement("w:r")
        new_text = OxmlElement("w:t")
        new_text.set(qn("xml:space"), "preserve")
        new_text.text = text
        new_run.append(new_text)
        for idx in range(end_idx, i - 1, -1):
            element.remove(children[idx])
        element.insert(i, new_run)
        children = list(element)
        i += 1


def replace_placeholders(
    source,
    values: dict[str, str],
) -> tuple[bytes, list[str]]:
    """Nahradí {{token}} i MERGEFIELD. Vrací (bytes, chybějící klíče)."""
    document = load_document(source)
    missing: list[str] = []
    seen_missing: set[str] = set()

    def substitute(text: str) -> str:
        def repl(match: re.Match[str]) -> str:
            key = match.group(1)
            if key not in values:
                if key not in seen_missing:
                    seen_missing.add(key)
                    missing.append(key)
                return ""
            raw = values.get(key) or ""
            if not raw and key not in seen_missing:
                seen_missing.add(key)
                missing.append(key)
            return _safe_text(raw)

        return TOKEN_RE.sub(repl, text or "")

    for paragraph in _iter_paragraphs(document):
        _replace_mergefields_in_paragraph(
            paragraph, values, missing, seen_missing
        )
        original = paragraph.text or ""
        if "{{" in original:
            _set_paragraph_text(paragraph, substitute(original))

    buffer = io.BytesIO()
    document.save(buffer)
    return strip_mail_merge_settings(buffer.getvalue()), missing


def build_placeholders_csv(scope: str) -> bytes:
    """CSV (UTF-8 BOM, ;) pro Word Korespondence – hlavičky = názvy polí."""
    items: list[PlaceholderDef] = placeholders_for_scope(scope)
    headers = [item.key for item in items]
    sample = [item.label for item in items]

    def escape(cell: str) -> str:
        if ";" in cell or '"' in cell or "\n" in cell:
            return '"' + cell.replace('"', '""') + '"'
        return cell

    lines = [
        ";".join(escape(h) for h in headers),
        ";".join(escape(s) for s in sample),
    ]
    content = "\r\n".join(lines) + "\r\n"
    return ("\ufeff" + content).encode("utf-8")
