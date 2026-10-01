"""Testy panelu Dokumenty (šablony, placeholdery, oprávnění)."""

from __future__ import annotations

import io
import tempfile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from docx import Document

from apps.accounts.roles import ROLE_ADMIN, ROLE_EMPLOYEE, ROLE_MANAGER, ROLE_READER
from apps.documents.merge import (
    append_placeholder_token,
    build_placeholders_csv,
    create_blank_docx_bytes,
    create_universal_template_bytes,
    extract_tokens_from_document,
    load_document,
    replace_placeholders,
    strip_mail_merge_settings,
    validate_placeholders_in_docx,
)
from apps.documents.models import DocumentTemplate, StoredDocument
from apps.documents.services import (
    append_placeholder_to_template,
    create_template,
    deactivate_template,
    delete_stored_document,
    ensure_universal_templates,
    generate_from_template,
    upload_template_version,
)
from apps.employees.models import Employee
from apps.technika.models import Vehicle

_MEDIA = tempfile.mkdtemp(prefix="mpdmb_docs_test_")


class DocumentsMergeTests(TestCase):
    def test_replace_mergefield_in_docx(self):
        data = create_blank_docx_bytes()
        data = append_placeholder_token(data, "zamestnanec.prijmeni")
        filled, missing = replace_placeholders(
            data, {"zamestnanec.prijmeni": "Novák", "dnes": "01.10.2026"}
        )
        doc = Document(io.BytesIO(filled))
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Novák", texts)
        self.assertNotIn("«zamestnanec.prijmeni»", texts)
        self.assertEqual(missing, [])

    def test_replace_curly_placeholder(self):
        buffer = io.BytesIO()
        doc = Document()
        doc.add_paragraph("{{zamestnanec.prijmeni}}")
        doc.save(buffer)
        filled, missing = replace_placeholders(
            buffer.getvalue(), {"zamestnanec.prijmeni": "Svoboda"}
        )
        out = Document(io.BytesIO(filled))
        texts = "\n".join(p.text for p in out.paragraphs)
        self.assertIn("Svoboda", texts)
        self.assertEqual(missing, [])

    def test_unknown_placeholder_rejected(self):
        buffer = io.BytesIO()
        doc = Document()
        doc.add_paragraph("{{neznamy.token}}")
        doc.save(buffer)
        unknown = validate_placeholders_in_docx(buffer.getvalue(), "zamestnanec")
        self.assertEqual(unknown, ["neznamy.token"])

    def test_empty_mergefield_listed_as_missing(self):
        data = append_placeholder_token(
            create_blank_docx_bytes(), "zamestnanec.telefon"
        )
        filled, missing = replace_placeholders(
            data, {"zamestnanec.telefon": ""}
        )
        self.assertIn("zamestnanec.telefon", missing)
        doc = Document(io.BytesIO(filled))
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("«zamestnanec.telefon»", texts)

    def test_universal_template_contains_mergefields(self):
        data = create_universal_template_bytes(
            "zamestnanec", title="Test"
        )
        doc = load_document(data)
        keys = extract_tokens_from_document(doc)
        self.assertIn("zamestnanec.prijmeni", keys)
        self.assertIn("dnes", keys)

    def test_strip_mail_merge_removes_sql_link(self):
        import zipfile
        from lxml import etree

        def settings_has_mail_merge(docx_bytes: bytes) -> bool:
            with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
                root = etree.fromstring(zf.read("word/settings.xml"))
            return any(
                etree.QName(node).localname == "mailMerge" for node in root
            )

        def with_mail_merge(docx_bytes: bytes) -> bytes:
            source = io.BytesIO(docx_bytes)
            output = io.BytesIO()
            with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
                output, "w"
            ) as zout:
                for info in zin.infolist():
                    data = zin.read(info.filename)
                    if info.filename == "word/settings.xml":
                        root = etree.fromstring(data)
                        mm = etree.SubElement(
                            root,
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}mailMerge",
                        )
                        etree.SubElement(
                            mm,
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}dataType",
                        ).set(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val",
                            "textFile",
                        )
                        data = etree.tostring(
                            root,
                            xml_declaration=True,
                            encoding="UTF-8",
                            standalone=True,
                        )
                    zout.writestr(info, data)
            return output.getvalue()

        tainted = with_mail_merge(create_blank_docx_bytes())
        self.assertTrue(settings_has_mail_merge(tainted))
        cleaned = strip_mail_merge_settings(tainted)
        self.assertFalse(settings_has_mail_merge(cleaned))
        filled, _ = replace_placeholders(tainted, {"dnes": "01.10.2026"})
        self.assertFalse(settings_has_mail_merge(filled))
        appended = append_placeholder_token(tainted, "zamestnanec.prijmeni")
        self.assertTrue(settings_has_mail_merge(appended))

    def test_universal_template_has_czech_label_and_comment(self):
        import zipfile

        data = create_universal_template_bytes(
            "zamestnanec", title="Test"
        )
        doc = load_document(data)
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("«Příjmení»", texts)
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            self.assertIn("word/comments.xml", names)
            comments = zf.read("word/comments.xml").decode("utf-8")
        self.assertIn("SYNERA", comments)
        self.assertTrue(
            "Příjmení" in comments or "prijmeni" in comments.lower()
        )

    def test_csv_has_bom_and_semicolon(self):
        raw = build_placeholders_csv("zamestnanec")
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        text = raw.decode("utf-8-sig")
        self.assertIn(";", text)
        self.assertIn("zamestnanec.prijmeni", text.splitlines()[0])


@override_settings(MEDIA_ROOT=_MEDIA)
class DocumentsAccessTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            username="doc_admin",
            password="pass-admin",
            is_staff=True,
            is_superuser=True,
        )
        Group.objects.get_or_create(name=ROLE_ADMIN)
        self.admin.groups.add(Group.objects.get(name=ROLE_ADMIN))

        self.manager = User.objects.create_user(
            username="doc_manager",
            password="pass-manager",
            is_staff=False,
        )
        Group.objects.get_or_create(name=ROLE_MANAGER)
        self.manager.groups.add(Group.objects.get(name=ROLE_MANAGER))

        self.employee_user = User.objects.create_user(
            username="doc_emp",
            password="pass-emp",
        )
        Group.objects.get_or_create(name=ROLE_EMPLOYEE)
        self.employee_user.groups.add(Group.objects.get(name=ROLE_EMPLOYEE))

        self.reader = User.objects.create_user(
            username="doc_reader",
            password="pass-reader",
        )
        Group.objects.get_or_create(name=ROLE_READER)
        self.reader.groups.add(Group.objects.get(name=ROLE_READER))

        self.employee = Employee.objects.create(
            first_name="Anna",
            last_name="Testova",
            internal_number="D001",
            is_active=True,
            user=self.employee_user,
            phone="777111222",
        )
        self.other = Employee.objects.create(
            first_name="Boris",
            last_name="Jiny",
            internal_number="D002",
            is_active=True,
        )
        self.vehicle = Vehicle.objects.create(
            plate="1A23456",
            name="Dodávka",
            responsible=self.employee,
            is_active=True,
        )
        self.client = Client()

    def test_admin_creates_template_and_generates(self):
        self.client.login(username="doc_admin", password="pass-admin")
        response = self.client.post(
            reverse("document_template_create"),
            {"name": "Potvrzení", "scope": "zamestnanec"},
        )
        self.assertEqual(response.status_code, 302)
        template = DocumentTemplate.objects.get(name="Potvrzení")
        append_placeholder_to_template(
            user=self.admin,
            template=template,
            placeholder_key="zamestnanec.prijmeni",
        )
        stored = generate_from_template(
            user=self.admin,
            template=template,
            employee=self.employee,
            title="Potvrzení Anna",
        )
        self.assertTrue(stored.file.name.endswith(".docx"))
        content = stored.file.read()
        doc = Document(io.BytesIO(content))
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Testova", texts)

        response = self.client.get(
            reverse("document_stored_download", args=[stored.pk])
        )
        self.assertEqual(response.status_code, 200)

    def test_manager_cannot_create_template(self):
        self.client.login(username="doc_manager", password="pass-manager")
        response = self.client.get(reverse("document_template_list"))
        self.assertEqual(response.status_code, 403)

    def test_manager_can_generate(self):
        template = create_template(
            user=self.admin,
            name="Pro vedoucího",
            scope="zamestnanec",
        )
        append_placeholder_to_template(
            user=self.admin,
            template=template,
            placeholder_key="zamestnanec.jmeno",
        )
        self.client.login(username="doc_manager", password="pass-manager")
        response = self.client.post(
            reverse("document_employee_generate", args=[self.employee.pk]),
            {
                "action": "generate",
                "template": template.pk,
                "title": "Vygenerováno vedoucím",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            StoredDocument.objects.filter(
                employee=self.employee, title="Vygenerováno vedoucím"
            ).exists()
        )

    def test_employee_sees_own_only(self):
        template = create_template(
            user=self.admin, name="Osobní", scope="zamestnanec"
        )
        own = generate_from_template(
            user=self.admin,
            template=template,
            employee=self.employee,
            title="Moje",
        )
        other_doc = generate_from_template(
            user=self.admin,
            template=template,
            employee=self.other,
            title="Cizí",
        )
        self.client.login(username="doc_emp", password="pass-emp")
        response = self.client.get(reverse("document_my_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Moje")
        self.assertNotContains(response, "Cizí")

        ok = self.client.get(reverse("document_stored_download", args=[own.pk]))
        self.assertEqual(ok.status_code, 200)
        forbidden = self.client.get(
            reverse("document_stored_download", args=[other_doc.pk])
        )
        self.assertEqual(forbidden.status_code, 403)

        templates = self.client.get(reverse("document_template_list"))
        self.assertEqual(templates.status_code, 403)

    def test_reader_blocked(self):
        self.client.login(username="doc_reader", password="pass-reader")
        response = self.client.get(reverse("documents_hub"))
        self.assertEqual(response.status_code, 403)

    def test_vehicle_generate(self):
        template = create_template(
            user=self.admin, name="Protokol vozidla", scope="vozidlo"
        )
        append_placeholder_to_template(
            user=self.admin,
            template=template,
            placeholder_key="vozidlo.spz",
        )
        stored = generate_from_template(
            user=self.admin,
            template=template,
            vehicle=self.vehicle,
            title="Protokol",
        )
        content = stored.file.read()
        doc = Document(io.BytesIO(content))
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("1A23456", texts)

    def test_upload_rejects_unknown_placeholder(self):
        buffer = io.BytesIO()
        doc = Document()
        doc.add_paragraph("{{foo.bar}}")
        doc.save(buffer)
        uploaded = SimpleUploadedFile(
            "bad.docx",
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
        self.client.login(username="doc_admin", password="pass-admin")
        response = self.client.post(
            reverse("document_template_create"),
            {
                "name": "Špatná",
                "scope": "zamestnanec",
                "file": uploaded,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Neznámé placeholdery")
        self.assertFalse(DocumentTemplate.objects.filter(name="Špatná").exists())

    def test_upload_template_version_from_uploaded_file(self):
        template = create_template(
            user=self.admin,
            name="Upload test",
            scope="zamestnanec",
        )
        data = create_blank_docx_bytes()
        data = append_placeholder_token(data, "zamestnanec.jmeno")
        uploaded = SimpleUploadedFile(
            "ok.docx",
            data,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
        version = upload_template_version(
            user=self.admin,
            template=template,
            uploaded_file=uploaded,
            notes="reupload",
        )
        self.assertEqual(version.version_number, 2)
        self.assertTrue(version.file.name.endswith(".docx"))

        self.client.login(username="doc_admin", password="pass-admin")
        uploaded2 = SimpleUploadedFile(
            "ok2.docx",
            data,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
        response = self.client.post(
            reverse("document_template_detail", args=[template.pk]),
            {
                "action": "version",
                "ver-notes": "via view",
                "ver-file": uploaded2,
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(template.versions.count(), 3)

    def test_upload_preserves_mail_merge(self):
        import zipfile
        from lxml import etree

        def settings_has_mail_merge(docx_bytes: bytes) -> bool:
            with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
                root = etree.fromstring(zf.read("word/settings.xml"))
            return any(
                etree.QName(node).localname == "mailMerge" for node in root
            )

        blank = create_blank_docx_bytes()
        source = io.BytesIO(blank)
        output = io.BytesIO()
        with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
            output, "w"
        ) as zout:
            for info in zin.infolist():
                raw = zin.read(info.filename)
                if info.filename == "word/settings.xml":
                    root = etree.fromstring(raw)
                    mm = etree.SubElement(
                        root,
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}mailMerge",
                    )
                    etree.SubElement(
                        mm,
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}dataType",
                    ).set(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val",
                        "textFile",
                    )
                    raw = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                zout.writestr(info, raw)
        tainted = output.getvalue()
        self.assertTrue(settings_has_mail_merge(tainted))
        template = create_template(
            user=self.admin,
            name="CSV keep",
            scope="zamestnanec",
        )
        uploaded = SimpleUploadedFile(
            "csv.docx",
            tainted,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
        version = upload_template_version(
            user=self.admin,
            template=template,
            uploaded_file=uploaded,
        )
        stored_bytes = version.file.read()
        self.assertTrue(settings_has_mail_merge(stored_bytes))

    def test_deactivate_template_and_delete_stored(self):
        template = create_template(
            user=self.admin,
            name="Ke smazání",
            scope="zamestnanec",
        )
        append_placeholder_to_template(
            user=self.admin,
            template=template,
            placeholder_key="zamestnanec.prijmeni",
        )
        stored = generate_from_template(
            user=self.admin,
            template=template,
            employee=self.employee,
            title="Doc ke smazání",
        )
        pk = stored.pk
        deactivate_template(user=self.admin, template=template)
        template.refresh_from_db()
        self.assertFalse(template.is_active)
        delete_stored_document(user=self.admin, document=stored)
        self.assertFalse(StoredDocument.objects.filter(pk=pk).exists())

        self.client.login(username="doc_admin", password="pass-admin")
        response = self.client.post(
            reverse("document_template_detail", args=[template.pk]),
            {"action": "activate"},
        )
        self.assertEqual(response.status_code, 302)
        template.refresh_from_db()
        self.assertTrue(template.is_active)

    def test_universal_templates_and_csv(self):
        ensure_universal_templates(user=self.admin, force_refresh=True)
        emp_tpl = DocumentTemplate.objects.get(name="Univerzální – zaměstnanec")
        self.assertTrue(emp_tpl.current_version)
        self.client.login(username="doc_admin", password="pass-admin")
        response = self.client.get(
            reverse("document_template_csv", args=[emp_tpl.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"zamestnanec.prijmeni", response.content)
        stored = generate_from_template(
            user=self.admin,
            template=emp_tpl,
            employee=self.employee,
            title="Uni fill",
        )
        content = stored.file.read()
        doc = Document(io.BytesIO(content))
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Testova", texts)
        self.assertIn("Anna", texts)
